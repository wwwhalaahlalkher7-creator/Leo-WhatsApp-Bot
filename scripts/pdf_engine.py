#!/usr/bin/env python3
import argparse, json, os, re, shutil, subprocess, sys, tempfile, zipfile
from pathlib import Path

try:
    import fitz
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_SECTION
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from pptx import Presentation
    from pptx.util import Inches as PptInches
    from pptx.enum.text import PP_ALIGN
except Exception as exc:
    print(f'engine import error: {exc}', file=sys.stderr)
    sys.exit(2)

IMAGE_EXTS = {'.jpg','.jpeg','.png','.webp','.bmp','.tif','.tiff'}
OFFICE_EXTS = {'.doc','.docx','.ppt','.pptx','.xls','.xlsx'}


def run(cmd):
    p = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if p.returncode:
        raise RuntimeError(p.stderr.strip() or 'command failed')
    return p


def page_count(pdf):
    with fitz.open(pdf) as doc:
        return len(doc)


def office_to_pdf(inp, outdir):
    profile = tempfile.mkdtemp(prefix='leo-lo-')
    try:
        run(['soffice','--headless','--convert-to','pdf','--outdir',str(outdir),'-env:UserInstallation=file://' + profile,str(inp)])
        out = Path(outdir) / (inp.stem + '.pdf')
        if not out.exists():
            candidates = list(Path(outdir).glob('*.pdf'))
            if not candidates: raise RuntimeError('LibreOffice did not create a PDF')
            out = candidates[0]
        return str(out)
    finally:
        shutil.rmtree(profile, ignore_errors=True)


def image_to_pdf(inputs, outdir):
    inputs = [Path(x) for x in inputs]
    if not inputs: raise RuntimeError('No images supplied')
    if len(inputs) > 80: raise RuntimeError('Maximum 80 images per PDF')
    out = Path(outdir) / ((inputs[0].stem if len(inputs) == 1 else 'images') + '.pdf')
    doc = fitz.open()
    try:
        for inp in inputs:
            if inp.suffix.lower() not in IMAGE_EXTS: raise RuntimeError(f'Unsupported image: {inp.name}')
            img = fitz.Pixmap(str(inp))
            if img.width <= 0 or img.height <= 0: raise RuntimeError(f'Invalid image: {inp.name}')
            if img.alpha: img = fitz.Pixmap(fitz.csRGB, img)
            page = doc.new_page(width=img.width, height=img.height)
            page.insert_image(fitz.Rect(0, 0, img.width, img.height), filename=str(inp))
        doc.save(str(out), garbage=4, deflate=True)
    finally:
        doc.close()
    return str(out)


def pdf_to_jpg(inp, outdir, max_pages):
    doc = fitz.open(inp)
    if len(doc) > max_pages: raise RuntimeError(f'PDF exceeds {max_pages} pages')
    paths=[]
    for i,p in enumerate(doc):
        pix = p.get_pixmap(matrix=fitz.Matrix(1.8,1.8), alpha=False)
        path = Path(outdir) / f'{inp.stem}-{i+1:03d}.jpg'
        pix.save(str(path), output='jpg', jpg_quality=92)
        paths.append(path)
    doc.close()
    if len(paths)==1: return str(paths[0])
    z=Path(outdir)/(inp.stem+'-images.zip')
    with zipfile.ZipFile(z,'w',zipfile.ZIP_DEFLATED) as f:
        for p in paths: f.write(p,p.name)
    return str(z)


def extract_pages(inp, max_pages):
    doc=fitz.open(inp)
    if len(doc)>max_pages: raise RuntimeError(f'PDF exceeds {max_pages} pages')
    pages=[]
    for p in doc:
        text=p.get_text('text').strip()
        pages.append((p.rect, text))
    doc.close(); return pages


def extract_text_result(inp, max_pages):
    pages = extract_pages(inp, max_pages)
    items = [{'page': i + 1, 'text': text} for i, (_, text) in enumerate(pages)]
    return {'pages': items, 'page_count': len(items), 'text': '\n\n'.join(f'[صفحة {x["page"]}]\n{x["text"]}' for x in items if x['text']).strip()}


def render_page(page, outpath, dpi=144):
    scale = dpi / 72.0
    pix = page.get_pixmap(matrix=fitz.Matrix(scale, scale), alpha=False)
    pix.save(str(outpath))
    return outpath


def pdf_to_word(inp,outdir,max_pages):
    # Fidelity-first conversion: every original PDF page is preserved visually.
    # A selectable/editable text appendix is added so the result remains useful.
    doc_pdf = fitz.open(inp)
    if len(doc_pdf)>max_pages: raise RuntimeError(f'PDF exceeds {max_pages} pages')
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(.25); sec.bottom_margin = Inches(.25); sec.left_margin = Inches(.25); sec.right_margin = Inches(.25)
    text_pages=[]
    for idx, page in enumerate(doc_pdf):
        if idx: doc.add_page_break()
        img_path = Path(outdir) / f'_word_page_{idx+1}.png'
        render_page(page, img_path, 144)
        available_w = sec.page_width - sec.left_margin - sec.right_margin
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        p.add_run().add_picture(str(img_path), width=available_w)
        text_pages.append(page.get_text('text').strip())
    # Keep text separately instead of corrupting the visual layout of the original pages.
    if any(text_pages):
        doc.add_page_break()
        h = doc.add_paragraph(); h.add_run('النص المستخرج من المستند').bold = True
        for i, text in enumerate(text_pages, 1):
            if i > 1: doc.add_page_break()
            p = doc.add_paragraph(); p.add_run(f'صفحة {i}').bold = True
            for line in text.splitlines():
                q = doc.add_paragraph(line)
                q.paragraph_format.space_after = Pt(2)
                for r in q.runs: r.font.size = Pt(10)
    out=Path(outdir)/(inp.stem+'.docx'); doc.save(str(out)); doc_pdf.close(); return str(out)


def _table_to_rows(table):
    rows=[]
    try:
        for row in table.extract():
            rows.append([str(c).strip() if c is not None else '' for c in row])
    except Exception:
        return []
    return rows


def pdf_to_excel(inp,outdir,max_pages):
    doc_pdf=fitz.open(inp)
    if len(doc_pdf)>max_pages: raise RuntimeError(f'PDF exceeds {max_pages} pages')
    wb=Workbook(); default=wb.active; wb.remove(default)
    any_table=False
    for pi,page in enumerate(doc_pdf,1):
        ws=wb.create_sheet(f'Page {pi}')
        ws.freeze_panes='A2'
        row=1
        # Prefer PyMuPDF's table detector for real tables.
        try:
            finder=page.find_tables()
            tables=getattr(finder,'tables',[]) or []
        except Exception:
            tables=[]
        if tables:
            for ti,table in enumerate(tables,1):
                rows=_table_to_rows(table)
                if not rows: continue
                any_table=True
                ws.cell(row,1,f'Table {ti}').font=Font(bold=True)
                row+=1
                for values in rows:
                    for col,val in enumerate(values,1):
                        c=ws.cell(row,col,val); c.alignment=Alignment(vertical='top',wrap_text=True)
                    row+=1
                row+=2
        else:
            # Fallback: preserve reading order and infer columns from text blocks.
            blocks=page.get_text('blocks')
            for b in sorted(blocks, key=lambda x:(round(x[1],1), x[0])):
                text=str(b[4] or '').strip()
                if not text: continue
                for line in text.splitlines():
                    cells=[x.strip() for x in re.split(r'\t+|\s{3,}', line) if x.strip()]
                    for col,val in enumerate(cells or [line.strip()],1):
                        ws.cell(row,col,val).alignment=Alignment(vertical='top',wrap_text=True)
                    row+=1
                row+=1
        if row == 1:
            ws.cell(1,1,'لا يوجد نص قابل للاستخراج في هذه الصفحة.')
        for col in range(1, min(ws.max_column, 40)+1):
            ws.column_dimensions[get_column_letter(col)].width=24
    # Always include a consolidated raw-text sheet for documents whose layout is complex.
    raw=wb.create_sheet('Extracted Text')
    raw['A1']='Page'; raw['B1']='Text'; raw['A1'].font=raw['B1'].font=Font(bold=True)
    r=2
    for pi,page in enumerate(doc_pdf,1):
        raw.cell(r,1,pi); raw.cell(r,2,page.get_text('text').strip()); raw.cell(r,2).alignment=Alignment(wrap_text=True,vertical='top'); r+=1
    raw.column_dimensions['A'].width=10; raw.column_dimensions['B'].width=100
    out=Path(outdir)/(inp.stem+'.xlsx'); wb.save(str(out)); doc_pdf.close(); return str(out)


def pdf_to_ppt(inp,outdir,max_pages):
    doc_pdf=fitz.open(inp)
    if len(doc_pdf)>max_pages: raise RuntimeError(f'PDF exceeds {max_pages} pages')
    first=doc_pdf[0].rect
    # Use a presentation ratio matching the source instead of forcing every PDF into 16:9.
    ratio=max(0.55,min(1.9, first.width/max(1,first.height)))
    if ratio >= 1.45: sw,sh=13.333,13.333/ratio
    elif ratio <= 1.0: sh=10.0; sw=sh*ratio
    else: sw,sh=12.0,12.0/ratio
    prs=Presentation(); prs.slide_width=PptInches(sw); prs.slide_height=PptInches(sh)
    blank=prs.slide_layouts[6]
    for i,page in enumerate(doc_pdf):
        pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False)
        img=Path(outdir)/f'_slide_{i+1}.png'; pix.save(str(img))
        slide=prs.slides.add_slide(blank)
        # Preserve aspect ratio and center the page; never stretch it.
        page_ratio=page.rect.width/max(1,page.rect.height)
        if page_ratio > ratio:
            w=prs.slide_width; h=int(w/page_ratio); x=0; y=int((prs.slide_height-h)/2)
        else:
            h=prs.slide_height; w=int(h*page_ratio); y=0; x=int((prs.slide_width-w)/2)
        slide.shapes.add_picture(str(img),x,y,width=w,height=h)
    out=Path(outdir)/(inp.stem+'.pptx'); prs.save(str(out)); doc_pdf.close(); return str(out)


def compress_pdf(inp,outdir):
    out=Path(outdir)/(inp.stem+'-compressed.pdf')
    profile=tempfile.mkdtemp(prefix='leo-gs-')
    try:
        run(['gs','-sDEVICE=pdfwrite','-dCompatibilityLevel=1.7','-dPDFSETTINGS=/ebook','-dNOPAUSE','-dQUIET','-dBATCH','-sOutputFile='+str(out),str(inp)])
        if out.stat().st_size >= Path(inp).stat().st_size: shutil.copy2(inp,out)
        return str(out)
    finally: shutil.rmtree(profile,ignore_errors=True)


def _translate_request(chunk, target):
    import urllib.parse, urllib.request, urllib.error
    headers={'User-Agent':'LeoBot/1.35.6 PDF Translator'}
    q=urllib.parse.quote(chunk, safe='')
    url=f'https://translate.googleapis.com/translate_a/single?client=gtx&sl=auto&tl={urllib.parse.quote(target)}&dt=t&q={q}'
    try:
        req=urllib.request.Request(url,headers=headers)
        with urllib.request.urlopen(req, timeout=35) as r:
            data=json.loads(r.read().decode('utf-8'))
        text=''.join(x[0] for x in (data[0] if isinstance(data,list) and data else []) if x and x[0]).strip()
        if text: return text
    except Exception:
        pass
    # Secondary public translator fallback.
    try:
        params=urllib.parse.urlencode({'q':chunk,'langpair':f'autodetect|{target}'})
        req=urllib.request.Request('https://api.mymemory.translated.net/get?'+params,headers=headers)
        with urllib.request.urlopen(req, timeout=35) as r:
            data=json.loads(r.read().decode('utf-8'))
        text=str(data.get('responseData',{}).get('translatedText') or '').strip()
        if text and not re.search(r'MYMEMORY WARNING|QUERY LENGTH LIMIT',text,re.I): return text
    except Exception:
        pass
    raise RuntimeError('Translation providers failed for a text segment')


def translate_text(text,target='ar'):
    if not text.strip(): return ''
    chunks=[]; cur=''
    # Keep request URLs comfortably below common proxy limits.
    for part in re.split(r'(\n\s*\n)', text):
        if len(cur)+len(part)>2800 and cur:
            chunks.append(cur); cur=''
        cur+=part
    if cur: chunks.append(cur)
    return ''.join(_translate_request(c,target) if c.strip() else c for c in chunks)


def _shape_rtl(text):
    try:
        import arabic_reshaper
        from bidi.algorithm import get_display
        if re.search(r'[\u0600-\u06ff]', text): return get_display(arabic_reshaper.reshape(text))
    except Exception:
        pass
    return text


def translate_pdf(inp,outdir,target='ar',max_pages=80):
    src=fitz.open(inp)
    if len(src)>max_pages: raise RuntimeError(f'PDF exceeds {max_pages} pages')
    doc=fitz.open()
    font='/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'
    for page in src:
        new=doc.new_page(width=page.rect.width,height=page.rect.height)
        # Start from a rendered copy so images, lines and original layout survive.
        pix=page.get_pixmap(matrix=fitz.Matrix(1.5,1.5),alpha=False)
        bg=Path(outdir)/f'_translation_bg_{page.number+1}.png'; pix.save(str(bg))
        new.insert_image(new.rect,filename=str(bg))
        blocks=page.get_text('blocks')
        translated_any=False
        for b in blocks:
            x0,y0,x1,y1,text=b[:5]
            text=str(text or '').strip()
            if not text: continue
            try:
                translated=translate_text(text,target)
            except Exception as exc:
                # Keep original text for this block instead of failing the whole document.
                translated=text
            translated=_shape_rtl(translated)
            # White out the original text area while retaining the rendered background.
            rect=fitz.Rect(x0,y0,x1,y1)
            new.draw_rect(rect,color=(1,1,1),fill=(1,1,1),overlay=True)
            size=max(7,min(18,float((y1-y0)*0.72)))
            rc=new.insert_textbox(rect,translated,fontfile=font,fontsize=size,color=(0,0,0),align=2 if target.lower() in ('ar','ara') else 0,overlay=True)
            if rc >= 0: translated_any=True
        if not translated_any and page.get_text('text').strip():
            # No usable block fit; keep the page visual rather than generating a blank translation.
            pass
    out=Path(outdir)/(Path(inp).stem+'-translated.pdf'); doc.save(str(out),garbage=4,deflate=True); doc.close(); src.close(); return str(out)


def main():
    ap=argparse.ArgumentParser()
    ap.add_argument('--operation',required=True); ap.add_argument('--input',action='append',required=True)
    ap.add_argument('--output-dir',required=True); ap.add_argument('--max-pages',type=int,default=80); ap.add_argument('--target',default='ar')
    a=ap.parse_args(); inputs=[Path(x) for x in a.input]; outdir=Path(a.output_dir); outdir.mkdir(parents=True,exist_ok=True)
    op=a.operation
    inp=inputs[0]
    if op=='extract-text':
        result=extract_text_result(inp,a.max_pages); print(json.dumps({'operation':op,**result},ensure_ascii=False)); return
    if op=='office-to-pdf': out=office_to_pdf(inp,outdir)
    elif op=='image-to-pdf': out=image_to_pdf(inputs,outdir)
    elif op=='pdf-to-jpg': out=pdf_to_jpg(inp,outdir,a.max_pages)
    elif op=='pdf-to-word': out=pdf_to_word(inp,outdir,a.max_pages)
    elif op=='pdf-to-excel': out=pdf_to_excel(inp,outdir,a.max_pages)
    elif op=='pdf-to-ppt': out=pdf_to_ppt(inp,outdir,a.max_pages)
    elif op=='compress-pdf': out=compress_pdf(inp,outdir)
    elif op=='translate-pdf': out=translate_pdf(inp,outdir,a.target,a.max_pages)
    else: raise RuntimeError('Unsupported operation')
    print(json.dumps({'output':out,'operation':op},ensure_ascii=False))

if __name__=='__main__':
    try: main()
    except Exception as e: print(str(e),file=sys.stderr); sys.exit(1)
